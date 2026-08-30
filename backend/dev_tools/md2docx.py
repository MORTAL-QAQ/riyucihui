# -*- coding: utf-8 -*-
"""将大创项目计划书 Markdown 转换为 Word（.docx）。

规则：
- '# ' 一级标题 / '## ' 二级标题
- '**（一）...**' 独立行 → 三级小标题（粗体段落）
- '| a | b |' 连续行 → 表格（第二行分隔线跳过）
- '- item' → 项目符号列表；'1. item' → 编号列表
- 其余 → 正文段落
"""
import io
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

# 命令行参数：python md2docx.py <源.md> <目标.docx>（缺省用计划书）
SRC = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\Administrator\Desktop\11\docs\大创项目计划书.md"
DST = sys.argv[2] if len(sys.argv) > 2 else re.sub(r"\.md$", ".docx", SRC)

doc = Document()

# 中文字体设置（宋体正文 / 黑体标题）
def set_cn_font(style, name="宋体", size=None, bold=None):
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold

set_cn_font(doc.styles["Normal"], size=11)
set_cn_font(doc.styles["Heading 1"], name="黑体", size=16, bold=True)
set_cn_font(doc.styles["Heading 2"], name="黑体", size=14, bold=True)
set_cn_font(doc.styles["Heading 3"], name="黑体", size=12, bold=True)

lines = io.open(SRC, encoding="utf-8").read().splitlines()

def is_table_sep(line):
    return bool(re.match(r"^\|[\s\-:|]+\|$", line))

def add_table(rows):
    # rows: list of list[str]（已去除 | 与空白）
    ncol = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=ncol)
    tbl.style = "Table Grid"
    for i, r in enumerate(rows):
        for j in range(ncol):
            cell = tbl.cell(i, j)
            cell.text = (r[j] if j < len(r) else "").strip()
            # 表头加粗
            if i == 0:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True

i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip():
        i += 1
        continue

    # 表格块
    if line.startswith("|"):
        block = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            ln = lines[i].strip()
            if not is_table_sep(ln):
                cells = [c.strip() for c in ln.strip("|").split("|")]
                block.append(cells)
            i += 1
        add_table(block)
        doc.add_paragraph()
        continue

    # 一级标题
    if line.startswith("# "):
        doc.add_heading(line[2:].strip(), level=1)
        i += 1
        continue
    # 二级标题
    if line.startswith("## "):
        doc.add_heading(line[3:].strip(), level=2)
        i += 1
        continue

    # 三级小标题：**（一）...**
    m = re.match(r"^\*\*（[一二三四五六七八九十]+）(.+?)\*\*$", line)
    if m:
        p = doc.add_paragraph()
        run = p.add_run("（" + m.group(1) + "）")
        run.bold = True
        run.font.size = Pt(12)
        i += 1
        continue

    # 项目符号列表
    if line.startswith("- "):
        doc.add_paragraph(line[2:].strip(), style="List Bullet")
        i += 1
        continue
    # 编号列表
    m = re.match(r"^(\d+)\.\s+(.*)$", line)
    if m:
        doc.add_paragraph(m.group(2).strip(), style="List Number")
        i += 1
        continue

    # 普通段落（含 **强调** 转粗体）
    p = doc.add_paragraph()
    parts = re.split(r"(\*\*.+?\*\*)", line)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part:
            p.add_run(part)
    i += 1

doc.save(DST)
print(f"已生成: {DST}")
